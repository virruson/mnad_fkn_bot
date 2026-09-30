#!/usr/bin/env python
"""
Загрузка дедлайнов, формул оценивания и модулей из Google Sheet или .xlsx.

Использование:
    python -m scripts.import_deadlines --xlsx scripts/deadlines_import.xlsx --dry-run
    python -m scripts.import_deadlines --xlsx scripts/deadlines_import.xlsx
    python -m scripts.import_deadlines --sheet-id <ID таблицы> [--dry-run]

ID таблицы можно задать переменной DEADLINES_SHEET_ID, путь к ключу сервисного
аккаунта — --credentials или GOOGLE_CREDENTIALS (по умолчанию scripts/credentials.json).
Таблицу нужно открыть на чтение для e-mail сервисного аккаунта.

Вкладки (первая строка — заголовки):
  «Дедлайны»  Предмет | Тип (ДЗ / Квиз / КВИЗ/КР / КР) | Название | Дата | Время
  «Формулы»   Предмет | Формула | Короткое название (необязательно)
  «Модули»    Название | Начало | Конец               (необязательная вкладка)

Даты — ДД.ММ.ГГГГ или ГГГГ-ММ-ДД, время — ЧЧ:ММ (пустое = 23:59).
Предмет ищется по точному названию из расписания; не найден — warning, строка пропускается.

Скрипт:
  - upsert дедлайнов по ключу предмет + тип + название (меняются дата и время)
  - upsert формул и коротких названий предметов, модулей по названию
  - НЕ удаляет ничего из БД
--dry-run — ничего не пишет в БД, только показывает, что было бы сделано
"""
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

# Запуск как модуль из корня проекта: python -m scripts.import_deadlines ...
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from bot.utils.database import SessionLocal
from bot.models import Subject, Deadline, Module

KINDS = {'дз': 'hw', 'квиз': 'quiz', 'квиз/кр': 'quiz', 'кр': 'test'}
TAB_DEADLINES, TAB_FORMULAS, TAB_MODULES = "Дедлайны", "Формулы", "Модули"
DEFAULT_CREDENTIALS = Path(__file__).resolve().parent / "credentials.json"


def fetch_tabs(sheet_id, credentials):
    """{название вкладки: строки без заголовка}. Отсутствующая вкладка → None."""
    import gspread
    spreadsheet = gspread.service_account(filename=credentials).open_by_key(sheet_id)
    tabs = {}
    for title in (TAB_DEADLINES, TAB_FORMULAS, TAB_MODULES):
        try:
            tabs[title] = spreadsheet.worksheet(title).get_all_values()[1:]
        except gspread.WorksheetNotFound:
            tabs[title] = None
    return tabs


def fetch_xlsx(path):
    """То же, что fetch_tabs, но из файла .xlsx с теми же вкладками."""
    from datetime import date, time
    from openpyxl import load_workbook

    def text(v):
        if isinstance(v, datetime):
            return v.strftime('%H:%M') if v.date() == date(1899, 12, 30) else v.strftime('%d.%m.%Y')
        if isinstance(v, date):
            return v.strftime('%d.%m.%Y')
        if isinstance(v, time):
            return v.strftime('%H:%M')
        return '' if v is None else str(v)

    wb = load_workbook(path, read_only=True, data_only=True)
    return {title: ([[text(v) for v in row] for row in wb[title].iter_rows(min_row=2, values_only=True)]
                    if title in wb.sheetnames else None)
            for title in (TAB_DEADLINES, TAB_FORMULAS, TAB_MODULES)}


def parse_date(value):
    value = value.strip()
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"дата «{value}» не в формате ДД.ММ.ГГГГ")


def parse_time(value):
    value = value.strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, '%H:%M').time()
    except ValueError:
        raise ValueError(f"время «{value}» не в формате ЧЧ:ММ")


def cells(row, count):
    """Строка таблицы → ровно count обрезанных ячеек (хвостовые пустые gspread не отдаёт)."""
    row = [str(c).strip() for c in row[:count]]
    return row + [''] * (count - len(row))


class Importer:
    def __init__(self, session):
        self.session = session
        self.subjects = {s.name.strip(): s for s in session.query(Subject).all()}
        self.stats = {'created': 0, 'updated': 0, 'unchanged': 0, 'skipped': 0}
        self.log = []

    def _record(self, action, text):
        self.stats[action] += 1
        if action != 'unchanged':
            mark = {'created': '+', 'updated': '~', 'skipped': 'WARNING'}[action]
            self.log.append(f"  {mark} {text}")

    def _warn(self, where, message):
        self._record('skipped', f"{where}: {message} — пропущено")

    def _subject(self, where, name):
        subject = self.subjects.get(name)
        if subject is None:
            self._warn(where, f"предмет «{name}» не найден")
        return subject

    @staticmethod
    def _apply(obj, values):
        """Записывает изменившиеся поля. Возвращает список «поле: было → стало»."""
        changes = []
        for field, value in values.items():
            old = getattr(obj, field)
            if old != value:
                changes.append(f"{field}: {old} → {value}")
                setattr(obj, field, value)
        return changes

    def import_modules(self, rows):
        for i, row in enumerate(rows, start=2):
            name, start, end = cells(row, 3)
            if not any((name, start, end)):
                continue
            where = f"{TAB_MODULES}!{i}"
            try:
                start_date, end_date = parse_date(start), parse_date(end)
            except ValueError as e:
                self._warn(where, e)
                continue
            if not name or start_date > end_date:
                self._warn(where, "нет названия или начало позже конца")
                continue
            module = self.session.query(Module).filter_by(name=name).first()
            if module is None:
                self.session.add(Module(name=name, start_date=start_date, end_date=end_date))
                self._record('created', f"модуль {name}: {start_date:%d.%m.%Y}–{end_date:%d.%m.%Y}")
            elif changes := self._apply(module, {'start_date': start_date, 'end_date': end_date}):
                self._record('updated', f"модуль {name}: {'; '.join(changes)}")
            else:
                self._record('unchanged', name)
            self.session.flush()

    def import_formulas(self, rows):
        for i, row in enumerate(rows, start=2):
            name, formula, short_name = cells(row, 3)
            if not name:
                continue
            subject = self._subject(f"{TAB_FORMULAS}!{i}", name)
            if subject is None:
                continue
            if formula.lower().startswith('итог'):   # «Итог = …» → только правая часть
                formula = formula.split('=', 1)[-1].strip()
            values = {'grading_formula': formula or subject.grading_formula,
                      'short_name': short_name or subject.short_name}
            if changes := self._apply(subject, values):
                self._record('updated', f"{name}: {'; '.join(changes)}")
            else:
                self._record('unchanged', name)

    def import_deadlines(self, rows):
        for i, row in enumerate(rows, start=2):
            name, kind_raw, title, date_raw, time_raw = cells(row, 5)
            if not any((name, kind_raw, title, date_raw)):
                continue
            where = f"{TAB_DEADLINES}!{i}"
            subject = self._subject(where, name)
            if subject is None:
                continue
            kind = KINDS.get(kind_raw.lower().replace(' ', ''))
            if kind is None:
                self._warn(where, f"тип «{kind_raw}» — ожидается ДЗ, Квиз, КВИЗ/КР или КР")
                continue
            if not title:
                self._warn(where, "нет названия")
                continue
            try:
                due_date, due_time = parse_date(date_raw), parse_time(time_raw)
            except ValueError as e:
                self._warn(where, e)
                continue

            label = f"{name} · {kind_raw} {title} · {due_date:%d.%m.%Y}{f' {due_time:%H:%M}' if due_time else ''}"
            deadline = self.session.query(Deadline).filter_by(
                subject_id=subject.id, kind=kind, title=title).first()
            if deadline is None:
                self.session.add(Deadline(subject_id=subject.id, kind=kind, title=title,
                                          due_date=due_date, due_time=due_time))
                self._record('created', label)
            elif changes := self._apply(deadline, {'due_date': due_date, 'due_time': due_time}):
                self._record('updated', f"{label} ({'; '.join(changes)})")
            else:
                self._record('unchanged', label)
            self.session.flush()  # чтобы повтор той же строки ниже нашёл уже добавленную запись


def run(tabs, dry_run=False):
    session = SessionLocal()
    importer = Importer(session)
    try:
        for title, method in ((TAB_MODULES, importer.import_modules),
                              (TAB_FORMULAS, importer.import_formulas),
                              (TAB_DEADLINES, importer.import_deadlines)):
            rows = tabs.get(title)
            if rows is None:
                print(f"[{title}] вкладка не найдена — пропускаю")
                continue
            before = len(importer.log)
            method(rows)
            print(f"[{title}] строк: {len(rows)}")
            print("\n".join(importer.log[before:]) or "  без изменений")

        if dry_run:
            session.rollback()
            print("[dry-run] изменения НЕ сохранены в БД")
        else:
            session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()
    return importer.stats


def main():
    parser = argparse.ArgumentParser(description="Загрузка дедлайнов и формул из Google Sheet или .xlsx")
    parser.add_argument('--xlsx', help="файл .xlsx с вкладками Дедлайны/Формулы/Модули вместо Google Sheet")
    parser.add_argument('--sheet-id', default=os.getenv('DEADLINES_SHEET_ID'),
                        help="ID Google-таблицы (или переменная DEADLINES_SHEET_ID)")
    parser.add_argument('--credentials', default=os.getenv('GOOGLE_CREDENTIALS', str(DEFAULT_CREDENTIALS)),
                        help="JSON-ключ сервисного аккаунта")
    parser.add_argument('--dry-run', action='store_true', help="не сохранять изменения в БД")
    args = parser.parse_args()
    if args.xlsx:
        tabs = fetch_xlsx(args.xlsx)
    elif args.sheet_id:
        tabs = fetch_tabs(args.sheet_id, args.credentials)
    else:
        parser.error("укажите --xlsx, --sheet-id или DEADLINES_SHEET_ID")

    stats = run(tabs, dry_run=args.dry_run)
    print(f"Создано: {stats['created']}, обновлено: {stats['updated']}, "
          f"без изменений: {stats['unchanged']}, пропущено: {stats['skipped']}")


if __name__ == '__main__':
    main()
