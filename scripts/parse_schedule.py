#!/usr/bin/env python
"""
Разбор вкладки модуля из xlsx-расписания программы в JSONL для scripts/import_schedule.py.

    python scripts/parse_schedule.py расписание.xlsx --sheet "6 модуль (предварительно)" \
        --links-from scripts/schedule_5module_parsed.jsonl --out scripts/schedule_6module_parsed.jsonl

Вкладка — два месячных блока (A–G и I–O): строка с числами дней, под ней строка с занятиями.
Ячейка — одно или несколько занятий через пустую строку:
    Предмет (Маголего)
    Преподаватель[, экзамен]
    Лекция
    18:10-19:30
    Семинар
    19:40-21:00
    ссылка
Ссылка на встречу: гиперссылка ячейки вида …/j/… (не запись), иначе — ссылка этого предмета
из --links-from (прошлые модули). Записи без времени («Сессия», экзамен без часа) пропускаются.
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

MONTHS = {'январь': 1, 'февраль': 2, 'март': 3, 'апрель': 4, 'май': 5, 'июнь': 6, 'июль': 7,
          'август': 8, 'сентябрь': 9, 'октябрь': 10, 'ноябрь': 11, 'декабрь': 12}
BLOCKS = (('A', 'G'), ('I', 'O'))
ELECTIVE = "(Маголего)"
TIME_RE = re.compile(r'^(\d{1,2}:\d{2})(?:\s*-\s*(\d{1,2}:\d{2}))?$')
LESSON_TYPES = ('Лекция', 'Семинар', 'Экзамен', 'Консультация')
SKIP_LINES = ('ссылка', 'запись')


def col_range(first, last):
    return [chr(c) for c in range(ord(first), ord(last) + 1)]


def month_of(ws, first_col):
    """«Ноябрь 2026» в строке 4 над блоком → (год, месяц)."""
    for row in range(1, 6):
        value = str(ws[f"{first_col}{row}"].value or '').strip().lower()
        parts = value.split()
        if len(parts) == 2 and parts[0] in MONTHS and parts[1].isdigit():
            return int(parts[1]), MONTHS[parts[0]]
    return None


def day_rows(ws, cols):
    """Строки, где в блоке стоят числа дней (5.0, 12.0 …)."""
    rows = []
    for r in range(1, ws.max_row + 1):
        nums = [ws[f"{c}{r}"].value for c in cols]
        if any(isinstance(v, (int, float)) for v in nums):
            rows.append(r)
    return rows


def parse_block(text):
    """Одно занятие-блок → [(subject, stream, teacher, lesson_type, start, end)]."""
    lines = [ln.strip() for ln in text.split('\n') if ln.strip()]
    if not lines:
        return []
    subject = lines[0]
    stream = None
    if ELECTIVE in subject:
        subject, stream = subject.replace(ELECTIVE, '').strip(), "Маголего"
    teacher, kind_hint = None, None
    rest = lines[1:]
    if rest and not TIME_RE.match(rest[0]) and rest[0].split()[0] not in LESSON_TYPES:
        teacher = rest.pop(0)
        if ',' in teacher:                          # «Мария Горденко, экзамен»
            teacher, hint = [p.strip() for p in teacher.split(',', 1)]
            kind_hint = 'Экзамен' if 'экзамен' in hint.lower() else None
    out, pending = [], kind_hint
    for line in rest:
        m = TIME_RE.match(line)
        if m:
            out.append((subject, stream, teacher, pending or 'Семинар', m.group(1).zfill(5),
                        m.group(2).zfill(5) if m.group(2) else None))
            pending = kind_hint
        elif line.split()[0] in LESSON_TYPES:
            pending = line.split()[0]
        elif line.lower() in SKIP_LINES:
            continue
    return out


def load_links(paths):
    links = {}
    for p in paths or []:
        for line in open(p, encoding='utf-8'):
            if line.strip():
                r = json.loads(line)
                if r.get('meeting_link'):
                    links[r['subject']] = r['meeting_link']
    return links


def parse_sheet(ws, links):
    rows = []
    for first, last in BLOCKS:
        cols = col_range(first, last)
        ym = month_of(ws, first)
        if not ym:
            continue
        for r in day_rows(ws, cols):
            year, month = ym
            prev = 0
            for c in cols:
                num = ws[f"{c}{r}"].value
                if not isinstance(num, (int, float)):
                    continue
                day = int(num)
                if day < prev:                      # 30, 1, 2 … — перешли в следующий месяц
                    month, year = (1, year + 1) if month == 12 else (month + 1, year)
                prev = day
                cell = ws[f"{c}{r + 1}"]
                if not cell.value:
                    continue
                cell_link = cell.hyperlink.target if cell.hyperlink else None
                if cell_link and '/j/' not in cell_link:
                    cell_link = None                # запись прошедшего занятия — не ссылка на встречу
                text = str(cell.value).replace('\r', '')
                for block in re.split(r'\n\s*\n', text):
                    for subject, stream, teacher, kind, start, end in parse_block(block):
                        rows.append({
                            "subject": subject, "stream": stream, "teacher": teacher,
                            "lesson_type": kind, "time_start": start, "time_end": end,
                            "meeting_link": cell_link or links.get(subject),
                            "raw_line": " | ".join(ln.strip() for ln in block.split('\n') if ln.strip()),
                            "date": date(year, month, day).isoformat(),
                            "cell": f"{c}{r + 1}", "sheet": ws.title,
                        })
    rows.sort(key=lambda x: (x['date'], x['time_start'], x['subject']))
    return rows


def main():
    ap = argparse.ArgumentParser(description="xlsx-вкладка модуля → JSONL для import_schedule")
    ap.add_argument('xlsx')
    ap.add_argument('--sheet', required=True)
    ap.add_argument('--links-from', action='append', help="JSONL прошлых модулей — ссылки по предмету")
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx)
    rows = parse_sheet(wb[args.sheet], load_links(args.links_from))
    with open(args.out, 'w', encoding='utf-8') as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    no_link = sorted({r['subject'] for r in rows if not r['meeting_link']})
    print(f"Занятий: {len(rows)} → {args.out}")
    if no_link:
        print("Без ссылки: " + ", ".join(no_link))


if __name__ == '__main__':
    main()
